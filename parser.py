"""
xlsx 파싱 모듈 v3.1: 설계표준안 Excel 파일에서 데이터 추출
- 체크박스(Form Control) 상태 읽기 ★ 신규
- 시트별 이미지 매핑 (행정도/GIS) ★ 신규
- Streamlit UploadedFile / 파일경로 / BytesIO 모두 지원
"""
import re
import zipfile
import io
import os
import warnings
from typing import Tuple, Dict, List
from models import CableInfo, EnclosureInfo, ImageInfo
from config import (
    HEADER_MAP, SUMMARY_MAP, BACKGROUND_CHECKS, FIELD_SURVEY,
    EXISTING_INFO, ROUTE_DESIGN, NEW_DESIGN, RM_CHECK, REGIONAL, COST_MAP,
    SNAPSHOT_SHEETS,
)

warnings.filterwarnings("ignore", category=UserWarning, module="openpyxl")


def safe_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


# ============================================================
# 파일 입력 / 워크북 열기
# ============================================================

def _get_file_bytes(uploaded_file) -> bytes:
    if isinstance(uploaded_file, str):
        with open(uploaded_file, 'rb') as f:
            return f.read()
    if isinstance(uploaded_file, bytes):
        return uploaded_file
    if hasattr(uploaded_file, 'read'):
        if hasattr(uploaded_file, 'seek'):
            try:
                uploaded_file.seek(0)
            except Exception:
                pass
        data = uploaded_file.read()
        if hasattr(uploaded_file, 'seek'):
            try:
                uploaded_file.seek(0)
            except Exception:
                pass
        if data and len(data) > 0:
            return data
    raise ValueError(f"파일을 읽을 수 없습니다. 타입: {type(uploaded_file).__name__}")


def _open_workbook(file_bytes: bytes):
    import openpyxl
    errors = []
    for mode_name, kwargs in [
        ("기본", dict(data_only=True, read_only=False)),
        ("읽기전용", dict(data_only=True, read_only=True)),
        ("수식", dict(data_only=False, read_only=False)),
    ]:
        try:
            return openpyxl.load_workbook(io.BytesIO(file_bytes), **kwargs)
        except Exception as e:
            errors.append(f"{mode_name}: {e}")

    # XML 복구 시도
    try:
        repaired = _repair_xlsx(file_bytes)
        if repaired:
            return openpyxl.load_workbook(io.BytesIO(repaired), data_only=True)
    except Exception as e:
        errors.append(f"XML복구: {e}")

    raise ValueError(
        "xlsx 파일을 열 수 없습니다.\n"
        + "\n".join(f"  {err}" for err in errors)
        + "\nExcel에서 '다른 이름으로 저장' 후 재시도하세요."
    )


def _repair_xlsx(file_bytes: bytes) -> bytes:
    try:
        inp = zipfile.ZipFile(io.BytesIO(file_bytes), 'r')
        buf = io.BytesIO()
        out = zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED)
        for name in inp.namelist():
            data = inp.read(name)
            if name.endswith('.xml') or name.endswith('.rels'):
                try:
                    text = data.decode('utf-8', errors='replace')
                    data = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text).encode('utf-8')
                except Exception:
                    pass
            out.writestr(name, data)
        out.close()
        inp.close()
        return buf.getvalue()
    except Exception:
        return None


# ============================================================
# 셀 읽기
# ============================================================

def read_cell(ws, col: str, row: int):
    try:
        return ws[f"{col}{row}"].value
    except Exception:
        return None


def parse_cell_map(ws, cell_map: dict) -> dict:
    return {key: read_cell(ws, col, row) for key, (col, row) in cell_map.items()}


# ============================================================
# ★ 체크박스(Form Control) 파싱 — 핵심 신규 기능
# ============================================================

def parse_checkboxes(file_bytes: bytes, target_sheet_name: str = "ENG시트(Checklist)") -> Dict[int, bool]:
    """
    xlsx ZIP에서 체크박스 상태를 추출.
    Excel Form Control은 ctrlProp XML에 checked="Checked" 속성으로 저장됨.
    
    Returns:
        {row_number(1-indexed): is_checked} 딕셔너리
    """
    checkbox_map = {}

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as z:
            names = z.namelist()

            # Step 1: 시트이름 → sheet?.xml 파일 번호 찾기
            sheet_file = _find_sheet_file(z, target_sheet_name)
            if not sheet_file:
                return checkbox_map

            sheet_num = re.search(r'sheet(\d+)\.xml', sheet_file)
            if not sheet_num:
                return checkbox_map
            sheet_idx = sheet_num.group(1)

            # Step 2: sheet?.xml.rels에서 rId → ctrlProp 매핑
            rels_path = f'xl/worksheets/_rels/sheet{sheet_idx}.xml.rels'
            if rels_path not in names:
                return checkbox_map

            rels_content = z.read(rels_path).decode('utf-8', errors='replace')
            rid_to_ctrl = {}
            for m in re.finditer(
                r'Id="(rId\d+)"[^>]*Target="\.\./?(ctrlProps/ctrlProp\d+\.xml)"',
                rels_content
            ):
                rid_to_ctrl[m.group(1)] = m.group(2)

            if not rid_to_ctrl:
                return checkbox_map

            # Step 3: sheet?.xml의 <controls> 섹션에서 rId → row 매핑
            sheet_content = z.read(f'xl/worksheets/sheet{sheet_idx}.xml').decode('utf-8', errors='replace')
            controls_match = re.search(r'<controls>(.*?)</controls>', sheet_content, re.DOTALL)
            if not controls_match:
                return checkbox_map

            rid_to_row = {}
            for m in re.finditer(
                r'r:id="(rId\d+)"[^>]*>.*?<xdr:row>(\d+)</xdr:row>',
                controls_match.group(1), re.DOTALL
            ):
                rid_to_row[m.group(1)] = int(m.group(2))  # 0-indexed

            # Step 4: ctrlProp XML에서 checked 상태 확인
            for rid, row_0 in rid_to_row.items():
                ctrl_path = rid_to_ctrl.get(rid)
                if not ctrl_path:
                    continue
                full_path = f'xl/{ctrl_path}'
                if full_path not in names:
                    continue
                try:
                    ctrl_content = z.read(full_path).decode('utf-8', errors='replace')
                    is_checked = 'checked="Checked"' in ctrl_content
                    row_1 = row_0 + 1  # 1-indexed로 변환
                    checkbox_map[row_1] = is_checked
                except Exception:
                    pass

    except Exception as e:
        print(f"체크박스 파싱 경고: {e}")

    return checkbox_map


def parse_checkboxes_from_cells(ws) -> Dict[int, bool]:
    """
    ★ 셀 기반 체크박스 파싱 (fallback).

    일부 양식은 ctrlProp XML 대신 C열 또는 D열에 직접 True/False (bool) 또는
    "TRUE"/"FALSE" (문자열)로 체크박스 상태를 저장합니다.
    CHECKBOX_ROW_MAP에 정의된 행만 읽어서 반환합니다.

    Returns:
        {row_number(1-indexed): is_checked} 딕셔너리
    """
    checkbox_map = {}

    # C열 또는 D열에서 체크박스 값 읽기 (양식에 따라 다름)
    # C열 Row 20이 "검토결과"이면 C열 사용, 아니면 D열도 시도
    check_cols = [3]  # C열 기본
    header_c20 = safe_str(ws.cell(row=20, column=3).value)
    if "검토결과" not in header_c20:
        check_cols.append(4)  # D열도 시도

    for row in CHECKBOX_ROW_MAP.keys():
        for col in check_cols:
            try:
                val = ws.cell(row=row, column=col).value
                if val is None:
                    if col == check_cols[-1]:  # 마지막 열에서도 None이면 False
                        checkbox_map[row] = False
                    continue
                elif isinstance(val, bool):
                    checkbox_map[row] = val
                    break
                elif isinstance(val, str):
                    v = val.strip().upper()
                    if v in ("TRUE", "1", "O", "○", "Y", "YES", "체크", "확인", "V", "√"):
                        checkbox_map[row] = True
                        break
                    elif v in ("FALSE", "0", "X", "×", "N", "NO", ""):
                        checkbox_map[row] = False
                        break
                    # 수량 문자열("6본", "1조")은 체크박스가 아니므로 무시 → 다음 열 시도
                    if col == check_cols[-1]:
                        pass  # 마지막 열에서도 매칭 안 되면 건너뜀
                elif isinstance(val, (int, float)):
                    checkbox_map[row] = bool(val)
                    break
            except Exception:
                pass

    return checkbox_map


# ============================================================
# ★ 시트별 이미지 매핑 — 핵심 신규 기능
# ============================================================

# 추출 대상 시트 (핵심 설계 도면)
TARGET_IMAGE_SHEETS = [
    "행정도(전)", "행정도(후)",
    "개황도(전)", "개황도(후)",  # ★ 양식 버전별 대체 이름
    "작업 전 GIS",
    "공문 및 사진",
    "선번도", "선번도(전)", "선번도(후)",
]

# 이미지 최소 크기 (아이콘/심볼 필터링)
MIN_IMAGE_SIZE = 30_000  # 30KB 이상만 의미 있는 이미지


# ============================================================
# ★ 시트 전체 스냅샷 렌더링 (Windows + Linux 크로스플랫폼)
# ============================================================

import platform
import subprocess
import shutil
import tempfile
import glob as _glob


def _find_libreoffice() -> str:
    """
    LibreOffice (soffice) 실행 파일 경로를 자동 탐색.
    Windows / Linux / macOS 모두 지원.
    """
    system = platform.system()

    if system == "Windows":
        candidates = [
            # 일반 설치
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
            # Chocolatey / Scoop
            os.path.expandvars(r"%PROGRAMFILES%\LibreOffice\program\soffice.exe"),
        ]
        # Glob으로 버전 폴더 탐색
        for pattern in [r"C:\Program Files\LibreOffice*\program\soffice.exe"]:
            candidates.extend(_glob.glob(pattern))
    elif system == "Darwin":  # macOS
        candidates = [
            "/Applications/LibreOffice.app/Contents/MacOS/soffice",
        ]
    else:  # Linux
        candidates = [
            "/usr/bin/libreoffice",
            "/usr/bin/soffice",
            "/usr/local/bin/libreoffice",
            "/snap/bin/libreoffice",
        ]

    # shutil.which 먼저 시도
    which_result = shutil.which("soffice") or shutil.which("libreoffice")
    if which_result:
        return which_result

    for path in candidates:
        if os.path.isfile(path):
            return path

    return ""


def _find_excel_com():
    """Windows에서 Excel COM 자동화 사용 가능 여부 확인"""
    if platform.system() != "Windows":
        return False
    try:
        import win32com.client
        return True
    except ImportError:
        return False


def render_sheet_snapshots(
    file_bytes: bytes,
    sheet_names: list = None,
    dpi: int = 150,
) -> Dict[str, ImageInfo]:
    """
    시트를 전체 스냅샷(PNG)으로 렌더링. 행정도의 선로, 심볼, 텍스트,
    현장사진이 모두 포함된 복합 이미지를 생성.

    렌더링 우선순위:
    1. LibreOffice (Windows / Linux / macOS)
    2. Excel COM 자동화 (Windows only)

    Returns:
        {시트이름: ImageInfo} — 시트별 전체 스냅샷 이미지
    """
    if sheet_names is None:
        sheet_names = SNAPSHOT_SHEETS

    sheet_index_map = _get_sheet_indices(file_bytes)
    if not sheet_index_map:
        return {}

    # 렌더링할 시트가 있는지 확인
    target_sheets = [s for s in sheet_names if s in sheet_index_map]
    if not target_sheets:
        return {}

    # 방법 1: LibreOffice
    lo_path = _find_libreoffice()
    if lo_path:
        result = _render_via_libreoffice(file_bytes, target_sheets, sheet_index_map, lo_path, dpi)
        if result:
            return result

    # 방법 2: Windows Excel COM
    if _find_excel_com():
        result = _render_via_excel_com(file_bytes, target_sheets, dpi)
        if result:
            return result

    return {}


def _get_sheet_indices(file_bytes: bytes) -> dict:
    """xlsx 내 시트이름 → 인덱스(0-based) 매핑 반환"""
    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as z:
            wb_xml = z.read('xl/workbook.xml').decode('utf-8')
            sheets = re.findall(r'<sheet[^>]*name="([^"]*)"', wb_xml)
            return {name: i for i, name in enumerate(sheets)}
    except Exception:
        return {}


def _render_via_libreoffice(
    file_bytes: bytes,
    target_sheets: list,
    sheet_index_map: dict,
    lo_path: str,
    dpi: int,
) -> Dict[str, ImageInfo]:
    """
    LibreOffice로 시트별 PDF → PNG 변환.
    매크로 대신 openpyxl로 대상 시트만 보이게 조작 후 --convert-to pdf 사용.
    """
    results = {}
    tmp_base = tempfile.mkdtemp(prefix="snap_lo_")

    try:
        for sheet_name in target_sheets:
            try:
                # (1) openpyxl로 대상 시트만 visible한 임시 xlsx 생성
                tmp_xlsx = os.path.join(tmp_base, f"snap_{sheet_name.replace('(','_').replace(')','_')}.xlsx")
                _create_single_sheet_xlsx(file_bytes, sheet_name, tmp_xlsx)

                # (2) LibreOffice --convert-to pdf
                pdf_dir = os.path.join(tmp_base, "pdf_out")
                os.makedirs(pdf_dir, exist_ok=True)

                cmd = [
                    lo_path, "--headless", "--calc",
                    "--convert-to", "pdf",
                    "--outdir", pdf_dir,
                    tmp_xlsx,
                ]
                proc = subprocess.run(
                    cmd, capture_output=True, text=True, timeout=120,
                    env={**os.environ, "HOME": tmp_base},  # 격리된 HOME (Linux 프로필 충돌 방지)
                )

                # (3) PDF 파일 찾기
                pdf_files = _glob.glob(os.path.join(pdf_dir, "*.pdf"))
                if not pdf_files:
                    continue

                pdf_path = pdf_files[0]
                png_data = _pdf_to_png(pdf_path, dpi)

                if png_data and len(png_data) > 5000:
                    results[sheet_name] = ImageInfo(
                        filename=f"snapshot_{sheet_name}.png",
                        data=png_data,

                        sheet_name=sheet_name,
                    )

                # 정리 (다음 시트를 위해)
                for f in pdf_files:
                    os.remove(f)
                os.remove(tmp_xlsx)

            except Exception:
                continue

    finally:
        shutil.rmtree(tmp_base, ignore_errors=True)

    return results


def _create_single_sheet_xlsx(file_bytes: bytes, target_sheet: str, output_path: str):
    """
    원본 xlsx에서 대상 시트만 남기고 나머지를 삭제한 사본 생성.
    openpyxl은 sheet_state='hidden'으로 해도 LibreOffice PDF 변환 시
    모든 시트가 포함되므로, 삭제 방식을 사용.
    """
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes))

    if target_sheet not in wb.sheetnames:
        wb.close()
        raise ValueError(f"시트 '{target_sheet}' 없음")

    # 대상 시트 외 모두 삭제
    sheets_to_remove = [name for name in wb.sheetnames if name != target_sheet]
    for name in sheets_to_remove:
        del wb[name]

    wb.save(output_path)
    wb.close()


def _render_via_excel_com(
    file_bytes: bytes,
    target_sheets: list,
    dpi: int,
) -> Dict[str, ImageInfo]:
    """
    Windows Excel COM 자동화를 사용하여 시트를 PDF로 내보내기.
    Excel이 설치된 Windows 환경에서만 동작.
    """
    try:
        import win32com.client
        import pythoncom
    except ImportError:
        return {}

    results = {}
    tmp_base = tempfile.mkdtemp(prefix="snap_excel_")

    try:
        tmp_xlsx = os.path.join(tmp_base, "source.xlsx")
        with open(tmp_xlsx, 'wb') as f:
            f.write(file_bytes)

        pythoncom.CoInitialize()
        excel = win32com.client.Dispatch("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False

        try:
            wb = excel.Workbooks.Open(os.path.abspath(tmp_xlsx))

            for sheet_name in target_sheets:
                try:
                    ws = wb.Worksheets(sheet_name)
                    pdf_path = os.path.join(tmp_base, f"{sheet_name}.pdf")

                    # Excel ExportAsFixedFormat (0 = PDF)
                    ws.ExportAsFixedFormat(
                        Type=0,
                        Filename=os.path.abspath(pdf_path),
                        Quality=0,  # standard
                    )

                    if os.path.exists(pdf_path):
                        png_data = _pdf_to_png(pdf_path, dpi)
                        if png_data and len(png_data) > 5000:
                            results[sheet_name] = ImageInfo(
                                filename=f"snapshot_{sheet_name}.png",
                                data=png_data,
        
                                sheet_name=sheet_name,
                            )
                except Exception:
                    continue

            wb.Close(False)
        finally:
            excel.Quit()
            pythoncom.CoUninitialize()

    except Exception:
        pass
    finally:
        shutil.rmtree(tmp_base, ignore_errors=True)

    return results


def _pdf_to_png(pdf_path: str, dpi: int = 150) -> bytes:
    """PDF 파일을 PNG 이미지 바이트로 변환 (다중 페이지 시 세로 결합)"""
    from PIL import Image as PILImage

    # 방법 1: PyMuPDF
    try:
        import fitz
        doc = fitz.open(pdf_path)
        if len(doc) == 0:
            doc.close()
            return None

        page_images = []
        for pn in range(len(doc)):
            page = doc[pn]
            mat = fitz.Matrix(dpi / 72, dpi / 72)
            pix = page.get_pixmap(matrix=mat)
            img = PILImage.open(io.BytesIO(pix.tobytes("png")))
            page_images.append(img)
        doc.close()

        return _combine_images(page_images)
    except ImportError:
        pass

    # 방법 2: pdf2image (poppler)
    try:
        from pdf2image import convert_from_path
        page_images = convert_from_path(pdf_path, dpi=dpi)
        if page_images:
            return _combine_images(page_images)
    except ImportError:
        pass

    # 방법 3: Pillow + subprocess (ghostscript)
    try:
        gs_cmd = shutil.which("gs") or shutil.which("gswin64c") or shutil.which("gswin32c")
        if gs_cmd:
            png_path = pdf_path.replace('.pdf', '.png')
            subprocess.run([
                gs_cmd, "-dNOPAUSE", "-dBATCH", "-sDEVICE=png16m",
                f"-r{dpi}", f"-sOutputFile={png_path}", pdf_path
            ], capture_output=True, timeout=60)
            if os.path.exists(png_path):
                with open(png_path, 'rb') as f:
                    return f.read()
    except Exception:
        pass

    return None


def _combine_images(page_images: list) -> bytes:
    """PIL Image 리스트를 세로로 결합하여 PNG 바이트 반환"""
    from PIL import Image as PILImage

    if not page_images:
        return None

    if len(page_images) == 1:
        final = page_images[0]
    else:
        total_h = sum(img.height for img in page_images)
        max_w = max(img.width for img in page_images)
        final = PILImage.new('RGB', (max_w, total_h), (255, 255, 255))
        y = 0
        for img in page_images:
            final.paste(img, (0, y))
            y += img.height

    buf = io.BytesIO()
    final.save(buf, format="PNG", optimize=True)
    return buf.getvalue()


def extract_images_by_sheet(file_bytes: bytes) -> Dict[str, List[ImageInfo]]:
    """
    시트별로 이미지를 매핑하여 추출.
    행정도(전/후), 작업 전 GIS 등 핵심 시트의 큰 이미지만 반환.
    ★ 동일 이미지가 여러 시트에 공유되면 content hash 기반으로 중복 표시.

    Returns:
        {"행정도(전)": [ImageInfo, ...], "행정도(후)": [...], ...}
    """
    import hashlib

    result = {}
    all_sheet_images = {}  # sheet_name → [(img_path, ImageInfo, hash)]

    try:
        with zipfile.ZipFile(io.BytesIO(file_bytes), 'r') as z:
            names = z.namelist()

            # Step 1: 시트이름 → sheet파일 → drawing 매핑
            sheet_drawings = {}
            for target_name in TARGET_IMAGE_SHEETS:
                sheet_file = _find_sheet_file(z, target_name)
                if not sheet_file:
                    continue

                sheet_num = re.search(r'sheet(\d+)\.xml', sheet_file)
                if not sheet_num:
                    continue

                rels_path = f'xl/worksheets/_rels/sheet{sheet_num.group(1)}.xml.rels'
                if rels_path not in names:
                    continue

                rels_content = z.read(rels_path).decode('utf-8', errors='replace')
                drawings = re.findall(
                    r'Target="\.\./?(drawings/drawing\d+\.xml)"',
                    rels_content
                )
                if drawings:
                    sheet_drawings[target_name] = drawings

            # Step 2: drawing → image 파일 매핑 (해시 포함)
            for sheet_name, drawings in sheet_drawings.items():
                sheet_images = []
                seen_images = set()

                for drawing_file in drawings:
                    drawing_rels = f'xl/drawings/_rels/{os.path.basename(drawing_file)}.rels'
                    if drawing_rels not in names:
                        continue

                    rels_content = z.read(drawing_rels).decode('utf-8', errors='replace')
                    image_paths = re.findall(
                        r'Target="\.\./?(media/image\d+\.\w+)"',
                        rels_content
                    )

                    for img_path in image_paths:
                        if img_path in seen_images:
                            continue
                        seen_images.add(img_path)

                        full_path = f'xl/{img_path}'
                        if full_path not in names:
                            continue

                        ext = img_path.rsplit('.', 1)[-1].lower()
                        if ext in ('emf', 'wmf', 'svg'):
                            continue

                        try:
                            info = z.getinfo(full_path)
                            if info.file_size < MIN_IMAGE_SIZE:
                                continue
                            if ext not in ('png', 'jpg', 'jpeg', 'bmp'):
                                continue

                            img_data = z.read(full_path)
                            img_hash = hashlib.md5(img_data).hexdigest()
                            img_info = ImageInfo(
                                filename=os.path.basename(img_path),
                                data=img_data,
                                description=f"{sheet_name}: {os.path.basename(img_path)} ({info.file_size:,}bytes)",
                                sheet_name=sheet_name,
                            )
                            sheet_images.append((img_path, img_info, img_hash))
                        except Exception:
                            pass

                if sheet_images:
                    all_sheet_images[sheet_name] = sheet_images

            # Step 3: ★ 크로스시트 중복 감지 — 행정도(전)/(후) 동일 이미지 처리
            global_hash_map = {}  # hash → first sheet_name
            for sheet_name, items in all_sheet_images.items():
                for img_path, img_info, img_hash in items:
                    if img_hash not in global_hash_map:
                        global_hash_map[img_hash] = sheet_name

            for sheet_name, items in all_sheet_images.items():
                unique_images = []
                for img_path, img_info, img_hash in items:
                    first_sheet = global_hash_map[img_hash]
                    if first_sheet != sheet_name:
                        # 이 이미지는 다른 시트에서 이미 등장 → 중복 표시
                        img_info.description += f" [공유:{first_sheet}]"
                        img_info._is_shared = True
                    else:
                        img_info._is_shared = False
                    unique_images.append(img_info)

                # 크기 내림차순 정렬
                unique_images.sort(key=lambda x: len(x.data), reverse=True)
                result[sheet_name] = unique_images

    except Exception as e:
        print(f"시트별 이미지 추출 경고: {e}")

    # ★ _shared_image_info: 행정도/개황도 전/후 동일여부 메타 정보
    before_key = "행정도(전)" if "행정도(전)" in result else ("개황도(전)" if "개황도(전)" in result else None)
    after_key = "행정도(후)" if "행정도(후)" in result else ("개황도(후)" if "개황도(후)" in result else None)
    if before_key and after_key:
        before_hashes = {hashlib.md5(img.data).hexdigest() for img in result[before_key]}
        after_hashes = {hashlib.md5(img.data).hexdigest() for img in result[after_key]}
        overlap = before_hashes & after_hashes
        total = before_hashes | after_hashes
        if total and len(overlap) / len(total) > 0.8:
            result["_images_shared_warning"] = (
                f"{before_key}과 {after_key}가 동일한 임베디드 이미지({len(overlap)}장)를 공유합니다. "
                "개별 이미지만으로는 전/후 차이를 확인할 수 없습니다. "
                "전체 스냅샷(LibreOffice/Excel) 또는 직접 이미지 업로드가 필요합니다."
            )

    return result


def extract_all_images(file_bytes: bytes) -> List[ImageInfo]:
    """하위 호환: 모든 이미지 평탄 추출 (기존 호환)"""
    sheet_images = extract_images_by_sheet(file_bytes)
    all_imgs = []
    for sheet_name, imgs in sheet_images.items():
        all_imgs.extend(imgs)
    return all_imgs


def _find_sheet_file(z, sheet_name: str) -> str:
    """시트이름으로 sheet?.xml 파일명 찾기"""
    try:
        wb_content = z.read('xl/workbook.xml').decode('utf-8', errors='replace')
        wb_rels = z.read('xl/_rels/workbook.xml.rels').decode('utf-8', errors='replace')

        # rId → sheet파일 매핑
        rid_map = {}
        for m in re.finditer(r'Id="(rId\d+)"[^>]*Target="(worksheets/sheet\d+\.xml)"', wb_rels):
            rid_map[m.group(1)] = m.group(2)

        # 시트이름 → rId 찾기
        # Pattern for: <sheet name="행정도(전)" sheetId="7" r:id="rId7"/>
        pattern = r'<sheet[^>]*name="([^"]*)"[^>]*r:id="(rId\d+)"'
        for m in re.finditer(pattern, wb_content):
            if m.group(1) == sheet_name:
                return rid_map.get(m.group(2), "")

        # 부분 매칭 (행정도 → 행정도(전))
        for m in re.finditer(pattern, wb_content):
            if sheet_name in m.group(1):
                return rid_map.get(m.group(2), "")

    except Exception:
        pass
    return ""


# ============================================================
# 케이블/함체 정보 파싱 (정규식)
# ============================================================

NETWORK_TIERS = r'기간망|간선망|지선망|가입자망|인입구간|인입망|국간'


def parse_cable_info(text: str) -> list:
    if not text:
        return []

    results = []

    # 패턴 1: [라벨 거리m] 망계위 =번호= 코어C/사용C (률%)
    pattern1 = (
        r'\[(.+?)\s+([\d,]+)\s*[mM]\]\s*'
        rf'({NETWORK_TIERS})\s*'
        r'(?:=\s*(\w+)\s*=\s*)?'
        r'(\d+)\s*C\s*/\s*(\d+)\s*C'
        r'\s*\(?\s*(\d+)\s*%\s*\)?'
    )
    for m in re.finditer(pattern1, text):
        results.append(CableInfo(
            label=m.group(1).strip(),
            length_m=int(m.group(2).replace(',', '')),
            network_tier=m.group(3),
            cable_id=m.group(4) or "",
            total_cores=int(m.group(5)),
            used_cores=int(m.group(6)),
            usage_rate=int(m.group(7)) / 100.0,
        ))

    # 패턴 2: [라벨 거리m] 망계위 종류 코어C/사용C(률%)
    if not results:
        pattern2 = (
            r'\[(.+?)\s+([\d,]+)\s*[mM]\]\s*'
            rf'({NETWORK_TIERS})\s*'
            r'(Dry|DRY|MSLT|SLT)?\s*'
            r'(\d+)\s*C\s*/\s*(\d+)\s*C'
            r'\s*\(?\s*(\d+)\s*%\s*\)?'
        )
        for m in re.finditer(pattern2, text):
            results.append(CableInfo(
                label=m.group(1).strip(),
                length_m=int(m.group(2).replace(',', '')),
                network_tier=m.group(3),
                cable_type=m.group(4) or "",
                total_cores=int(m.group(5)),
                used_cores=int(m.group(6)),
                usage_rate=int(m.group(7)) / 100.0,
            ))

    # 패턴 3: 거리 없는 케이블
    pattern3 = (
        r'\[(.+?)\]\s*'
        rf'({NETWORK_TIERS})\s*'
        r'(Dry|DRY|MSLT|SLT)?\s*'
        r'(\d+)\s*C\s*/\s*(\d+)\s*C'
        r'\s*\(?\s*(\d+)\s*%\s*\)?'
    )
    for m in re.finditer(pattern3, text):
        label = m.group(1).strip()
        if any(c.label == label for c in results):
            continue
        if re.search(r'\d+\s*[mM]', label):
            continue
        results.append(CableInfo(
            label=label, length_m=0,
            network_tier=m.group(2),
            cable_type=m.group(3) or "",
            total_cores=int(m.group(4)),
            used_cores=int(m.group(5)),
            usage_rate=int(m.group(6)) / 100.0,
        ))

    # 패턴 4: 가장 단순한 형태
    if not results:
        pattern4 = (
            rf'({NETWORK_TIERS})\s*'
            r'(?:=?\s*(\w+)\s*=?\s*)?'
            r'(\d+)\s*C\s*/\s*(\d+)\s*C'
            r'\s*\(?\s*(\d+)\s*%\s*\)?'
        )
        for m in re.finditer(pattern4, text):
            results.append(CableInfo(
                label="케이블", network_tier=m.group(1),
                cable_id=m.group(2) or "",
                total_cores=int(m.group(3)),
                used_cores=int(m.group(4)),
                usage_rate=int(m.group(5)) / 100.0,
            ))

    return results


def parse_enclosure_info(text: str) -> list:
    if not text:
        return []

    results = []

    # 패턴 1: [라벨] 전산번호 / SK번호 / 접속 NC
    for m in re.finditer(r'\[(.+?)\]\s*(\S+)\s*/\s*(\S+)\s*/\s*접속\s*(\d+)\s*C', text):
        results.append(EnclosureInfo(
            label=m.group(1).strip(),
            pole_id=m.group(2).strip(),
            sk_name=m.group(3).strip(),
            splice_cores=int(m.group(4)),
        ))

    if not results:
        for m in re.finditer(r'\[(.+?)\]\s*(\S+)\s*/?접속\s*(\d+)\s*C', text):
            results.append(EnclosureInfo(
                label=m.group(1).strip(),
                pole_id=m.group(2).strip(),
                splice_cores=int(m.group(3)),
            ))

    return results


# ============================================================
# 시트 이름 찾기
# ============================================================

def find_eng_sheet(wb) -> str:
    for name in ['ENG시트(Checklist)', 'ENG시트', 'Checklist', 'ENG']:
        if name in wb.sheetnames:
            return name
    for sheet in wb.sheetnames:
        if 'ENG' in sheet.upper() and '예시' not in sheet:
            return sheet
    for sheet in wb.sheetnames:
        if 'CHECK' in sheet.upper():
            return sheet
    return wb.sheetnames[0]


def find_cost_sheet(wb) -> str:
    for sheet in wb.sheetnames:
        if '원가' in sheet:
            return sheet
    return ""


# ============================================================
# 체크박스 → 항목명 매핑
# ============================================================

# ENG시트 체크박스가 위치한 행 → 항목명 매핑
CHECKBOX_ROW_MAP = {
    21: "이설요청_영배시스템",  # Row 21 = 영배시스템 등록 체크박스 (E열 라벨: '영업배전시스템')
    22: "이설요청_한전",        # Row 22 = 한전 이설요청서 첨부 체크박스 (E열 라벨: '이설요청서')
    23: "이설요청_지자체",
    24: "이설요청_공공기관",
    25: "이설요청_기타",
    26: "공사유형_가",  # 한전주 이설
    27: "공사유형_나",  # 지중화
    28: "공사유형_다",  # 도로확장 등
    29: "공사유형_라",  # 원인자
    30: "공사유형_마",  # 그 외
    31: "병행공사",
    32: "원인자판정조서",
    35: "GIS_불일치_관로전주",
    36: "GIS_불일치_케이블함체",
    37: "GIS_기타특이",
    38: "자가주관로",
    39: "타사주관로",
    40: "자가주인허가",
    41: "폭탄함체",
    50: "기설통신주관로",
    51: "병행관로전주",
    52: "단독관로전주",
    53: "한전주이설",
    57: "TB공동투자",
    58: "함체신설",
    59: "6차선횡단",
    60: "임의횡단_종말주_지상고",
    61: "한전위해개소",
    62: "코아링케이블링",
    63: "재활용케이블",
    64: "코아접속수량",
    65: "회선분산설계",
    66: "추가기준",
}


def map_checkboxes_to_items(checkbox_states: Dict[int, bool]) -> Dict[str, bool]:
    """체크박스 row → 항목명으로 변환"""
    result = {}
    for row, name in CHECKBOX_ROW_MAP.items():
        result[name] = checkbox_states.get(row, False)
    return result


# ============================================================
# 추가 데이터 수집
# ============================================================

def _parse_additional_data(ws, data: dict):
    for key, col, row in [
        # 설계 이설루트 상세 (D열: 상세 사유)
        ("기설루트사용_상세", 'D', 50),
        ("병행관로전주_상세", 'D', 51),
        ("단독관로전주_상세", 'D', 52),
        ("한전주이설_상세", 'D', 53),
        # 설계 신설정보 (C열: 수량, D열: 상세)
        ("절체이설_수량", 'C', 54),
        ("절체이설_상세", 'D', 54),
        ("용량증설_수량", 'C', 55),
        ("용량증설_상세", 'D', 55),
        ("다대화_수량", 'C', 56),
        ("다대화_상세", 'D', 56),
        ("TB공동투자_상세", 'D', 57),
        ("함체신설_상세", 'D', 58),
    ]:
        try:
            v = safe_str(read_cell(ws, col, row))
            if v and not data.get(key):
                data[key] = v
        except Exception:
            pass


# ============================================================
# 공사비 교차 검증
# ============================================================

def _crosscheck_cost(data: dict, wb, cost_name: str):
    """
    ENG시트 D7 (공사비)은 '=원가계산서!I40' 수식.
    openpyxl data_only=True는 캐시된 값만 반환하므로,
    파일이 Excel 외부에서 수정되면 캐시가 오래될 수 있음.
    
    원가계산서의 총계(I40)와 비교하여 불일치 시 보정.
    """
    eng_cost = data.get("공사비")
    cost_총계 = data.get("총계")

    def _to_num(v):
        if v is None:
            return None
        try:
            n = float(str(v).replace(',', '').replace('원', '').strip())
            if n > 0:
                return n
        except (ValueError, TypeError):
            pass
        return None

    eng_val = _to_num(eng_cost)
    총계_val = _to_num(cost_총계)

    # Case 1: ENG 공사비가 수식 문자열이거나 비정상 (=원가계산서!I40)
    if eng_val is None and 총계_val is not None:
        data["공사비"] = int(총계_val)
        data["_공사비_source"] = "원가계산서 총계 (ENG시트 수식 미해석)"
        return

    # Case 2: 둘 다 존재하고 일치 → 정상
    if eng_val is not None and 총계_val is not None:
        if abs(eng_val - 총계_val) < 10:  # 반올림 오차 허용
            return  # 정상

        # Case 3: 불일치 → 원가계산서 총계 우선 (더 신뢰)
        data["공사비"] = int(총계_val)
        data["_공사비_source"] = f"원가계산서 총계 (ENG={eng_val:,.0f} ≠ 총계={총계_val:,.0f})"
        return

    # Case 4: ENG만 있으면 그대로 사용
    if eng_val is not None:
        return

    # Case 5: 둘 다 없으면 직접 셀 읽기 시도
    if cost_name and cost_name in wb.sheetnames:
        try:
            ws_cost = wb[cost_name]
            v = read_cell(ws_cost, 'I', 40)
            val = _to_num(v)
            if val:
                data["공사비"] = int(val)
                data["_공사비_source"] = "원가계산서 I40 직접 읽기"
        except Exception:
            pass


# ============================================================
# 공문 및 사진 시트 확인
# ============================================================

def check_doc_photo_sheet(wb) -> Dict[str, bool]:
    """공문 및 사진 시트에 콘텐츠가 있는지 확인"""
    result = {"has_content": False, "has_images": False, "title": ""}

    for sheet_name in wb.sheetnames:
        if '공문' in sheet_name and '사진' in sheet_name:
            ws = wb[sheet_name]
            # 제목 확인
            title_val = read_cell(ws, 'B', 1)
            if title_val:
                result["title"] = safe_str(title_val)
                result["has_content"] = True

            # 데이터 존재 확인 (몇개 셀만 체크)
            for row in range(2, min(20, ws.max_row + 1)):
                for col in ['A', 'B', 'C', 'D']:
                    v = read_cell(ws, col, row)
                    if v and safe_str(v) not in ("", "None"):
                        result["has_content"] = True
                        break

            break

    return result


# ============================================================
# 메인 파싱 함수
# ============================================================

def parse_design_xlsx(uploaded_file) -> Tuple[dict, Dict[str, List[ImageInfo]], Dict[str, ImageInfo]]:
    """
    설계표준안 xlsx 파싱.

    Returns:
        (data_dict, sheet_images_dict, snapshots_dict)
        - data_dict: 파싱된 데이터 (체크박스 상태 포함)
        - sheet_images_dict: {"행정도(전)": [ImageInfo,...], ...} — 개별 이미지
        - snapshots_dict: {"행정도(전)": ImageInfo, ...} — ★ 시트 전체 스냅샷
    """
    file_bytes = _get_file_bytes(uploaded_file)

    if len(file_bytes) < 100:
        raise ValueError(f"파일 크기가 너무 작습니다 ({len(file_bytes)} bytes)")

    wb = _open_workbook(file_bytes)
    data = {}

    try:
        # 1. ENG시트 셀 데이터 파싱
        eng_name = find_eng_sheet(wb)
        ws = wb[eng_name]

        data.update(parse_cell_map(ws, HEADER_MAP))
        data.update(parse_cell_map(ws, SUMMARY_MAP))
        data.update(parse_cell_map(ws, BACKGROUND_CHECKS))
        data.update(parse_cell_map(ws, FIELD_SURVEY))
        data.update(parse_cell_map(ws, EXISTING_INFO))
        data.update(parse_cell_map(ws, ROUTE_DESIGN))
        data.update(parse_cell_map(ws, NEW_DESIGN))
        data.update(parse_cell_map(ws, RM_CHECK))
        data.update(parse_cell_map(ws, REGIONAL))
        _parse_additional_data(ws, data)

        # 2. ★ 체크박스 상태 파싱 (ctrlProp XML + 셀 bool 병합)
        xml_checkboxes = parse_checkboxes(file_bytes, eng_name)
        cell_checkboxes = parse_checkboxes_from_cells(ws)

        # ★ 두 방식 결과를 병합 (어느 한쪽이라도 True면 True)
        if xml_checkboxes and cell_checkboxes:
            raw_checkboxes = {}
            all_rows = set(xml_checkboxes.keys()) | set(cell_checkboxes.keys())
            for row in all_rows:
                xml_val = xml_checkboxes.get(row, False)
                cell_val = cell_checkboxes.get(row, False)
                raw_checkboxes[row] = xml_val or cell_val
            checkbox_method = "merged(ctrlProp+cell)"
        elif xml_checkboxes:
            raw_checkboxes = xml_checkboxes
            checkbox_method = "ctrlProp"
        elif cell_checkboxes:
            raw_checkboxes = cell_checkboxes
            checkbox_method = "cell_boolean"
        else:
            raw_checkboxes = {}
            checkbox_method = "none"

        checkbox_items = map_checkboxes_to_items(raw_checkboxes)
        data["_checkboxes"] = checkbox_items
        data["_checkbox_count"] = sum(1 for v in checkbox_items.values() if v)
        data["_checkbox_total"] = len(checkbox_items)
        data["_checkbox_method"] = checkbox_method

        # 3. 원가계산서 파싱
        cost_name = find_cost_sheet(wb)
        if cost_name:
            try:
                data.update(parse_cell_map(wb[cost_name], COST_MAP))
            except Exception as e:
                data["_cost_error"] = str(e)

        # 3-1. ★ 공사비 교차검증
        _crosscheck_cost(data, wb, cost_name)

        # 4. 공문 및 사진 시트 확인
        doc_info = check_doc_photo_sheet(wb)
        data["_doc_photo_has_content"] = doc_info["has_content"]

        # 5. 케이블/함체 파싱
        data["_parsed_cables"] = parse_cable_info(safe_str(data.get("케이블정보_상세")))
        data["_parsed_old_enclosures"] = parse_enclosure_info(safe_str(data.get("함체정보_상세")))
        data["_parsed_new_enclosures"] = parse_enclosure_info(safe_str(data.get("함체신설_상세")))
        data["_parsed_new_cables"] = parse_cable_info(safe_str(data.get("다대화_상세")))
        data["_parsed_upgrade_cables"] = parse_cable_info(safe_str(data.get("용량증설_상세")))
        data["_parsed_splice_cables"] = parse_cable_info(safe_str(data.get("절체이설_상세")))

        data["_sheet_names"] = wb.sheetnames

    except Exception as e:
        data["_parse_error"] = str(e)
        data["_sheet_names"] = getattr(wb, 'sheetnames', [])
    finally:
        try:
            wb.close()
        except Exception:
            pass

    # 6. ★ 시트별 개별 이미지 추출 (행정도/GIS 중심)
    sheet_images = extract_images_by_sheet(file_bytes)

    # 7. ★★ 시트 전체 스냅샷 렌더링 (LibreOffice)
    snapshots = {}
    try:
        snapshots = render_sheet_snapshots(file_bytes)
    except Exception as e:
        data["_snapshot_error"] = str(e)

    return data, sheet_images, snapshots


# ============================================================
# 하위호환: 기존 코드에서 (data, flat_images) 형태로 사용하는 경우
# ============================================================

def parse_design_xlsx_flat(uploaded_file) -> Tuple[dict, list]:
    """하위호환용: 이미지를 flat list로 반환"""
    data, sheet_images, _snapshots = parse_design_xlsx(uploaded_file)
    flat = []
    for imgs in sheet_images.values():
        flat.extend(imgs)
    return data, flat


# ============================================================
# 포맷팅 유틸
# ============================================================

def format_parsed_summary(data: dict) -> str:
    lines = []
    lines.append("=" * 50)
    lines.append("📋 설계서 파싱 결과 요약")
    lines.append("=" * 50)

    lines.append(f"\n▶ 공사명: {safe_str(data.get('공사명'))}")
    lines.append(f"▶ 사업구분: {safe_str(data.get('사업구분'))}")
    lines.append(f"▶ 공사방안: {safe_str(data.get('공사방안'))}")
    lines.append(f"▶ 공사유형: {safe_str(data.get('공사유형'))}")
    lines.append(f"▶ 요청주체: {safe_str(data.get('요청주체'))}")
    lines.append(f"▶ 현장주소: {safe_str(data.get('현장주소'))}")

    공사비 = data.get('공사비')
    try:
        lines.append(f"▶ 공사비: {int(float(공사비)):,}원")
    except (TypeError, ValueError):
        lines.append(f"▶ 공사비: {safe_str(공사비)}")

    # 체크박스 상태
    cb = data.get("_checkboxes", {})
    if cb:
        checked = sum(1 for v in cb.values() if v)
        lines.append(f"\n--- 체크박스 ({checked}/{len(cb)} 체크) ---")
        for name, val in cb.items():
            mark = "✅" if val else "⬜"
            lines.append(f"  {mark} {name}")

    lines.append(f"\n--- 설계 요약 ---")
    lines.append(f"이설방법: {safe_str(data.get('이설방법'))}")
    lines.append(f"절체사유: {safe_str(data.get('절체사유'))}")
    lines.append(f"케이블 신설: {safe_str(data.get('케이블_신설'))}")
    lines.append(f"케이블 철거: {safe_str(data.get('케이블_철거'))}")

    lines.append(f"\n--- 기설정보 ---")
    lines.append(f"전주: {safe_str(data.get('전주정보_수량'))} — {safe_str(data.get('전주정보_상세'))}")
    lines.append(f"케이블: {safe_str(data.get('케이블정보_수량'))} — {safe_str(data.get('케이블정보_상세'))}")
    lines.append(f"함체: {safe_str(data.get('함체정보_수량'))} — {safe_str(data.get('함체정보_상세'))}")

    cables = data.get("_parsed_cables", [])
    if cables:
        lines.append(f"\n--- 케이블 ({len(cables)}건) ---")
        for c in cables:
            lines.append(f"  {c.label}: {c.network_tier} {c.total_cores}C/{c.used_cores}C ({c.usage_rate:.0%}) {c.length_m}m")

    if data.get("_parse_error"):
        lines.append(f"\n⚠️ 파싱 경고: {data['_parse_error']}")

    return "\n".join(lines)
